"""
Tests for T2-E — PPTX / XLSX / HTML ingestion.

Validates:
- get_pptx_pages() returns one chunk per non-empty slide with slide number.
- get_pptx_pages() skips empty slides.
- get_pptx_pages() degrades gracefully when python-pptx is absent.
- get_xlsx_chunks() prefixes every chunk with [Sheet: name] + header.
- get_xlsx_chunks() handles multiple sheets and batches rows correctly.
- get_xlsx_chunks() handles header-only sheets.
- get_xlsx_chunks() degrades gracefully when openpyxl is absent.
- get_html_pages() extracts clean text via trafilatura (mocked).
- get_html_pages() falls back to stdlib html.parser on trafilatura failure.
- get_html_pages() strips script/style tags in fallback mode.
- process_single_document() assigns correct doc_type for each new extension.
- _ALLOWED_EXTENSIONS includes all new types in both ingest and router.
"""

import os
import sys
import tempfile
from unittest.mock import MagicMock, patch

import pytest


# ─────────────────────────────────────────────────────────────────────────────
# PPTX
# ─────────────────────────────────────────────────────────────────────────────

class TestGetPptxPages:
    def _make_slide(self, *texts):
        """Return a fake pptx slide whose shapes yield the given texts."""
        shapes = []
        for t in texts:
            shape = MagicMock()
            shape.has_text_frame = True
            shape.text_frame.text = t
            shapes.append(shape)
        slide = MagicMock()
        slide.shapes = shapes
        return slide

    def _make_presentation(self, slides):
        prs = MagicMock()
        prs.slides = slides
        return prs

    def test_returns_one_chunk_per_slide(self):
        from backend.services.ingest import get_pptx_pages

        slides = [
            self._make_slide("Title One", "Body of slide one"),
            self._make_slide("Title Two", "Body of slide two"),
        ]
        fake_prs = self._make_presentation(slides)

        with patch("backend.services.ingest.Presentation", return_value=fake_prs, create=True):
            # We need pptx importable; patch the import
            import types
            fake_pptx_module = types.SimpleNamespace(Presentation=lambda p: fake_prs)
            with patch.dict(sys.modules, {"pptx": fake_pptx_module}):
                pages = get_pptx_pages("deck.pptx")

        assert len(pages) == 2
        assert pages[0][1] == 1
        assert pages[1][1] == 2

    def test_skips_empty_slides(self):
        from backend.services.ingest import get_pptx_pages

        empty_shape = MagicMock()
        empty_shape.has_text_frame = True
        empty_shape.text_frame.text = "   "

        empty_slide = MagicMock()
        empty_slide.shapes = [empty_shape]

        full_slide = self._make_slide("Real content here")
        fake_prs = self._make_presentation([empty_slide, full_slide])

        import types
        fake_pptx_module = types.SimpleNamespace(Presentation=lambda p: fake_prs)
        with patch.dict(sys.modules, {"pptx": fake_pptx_module}):
            pages = get_pptx_pages("deck.pptx")

        assert len(pages) == 1
        assert "Real content" in pages[0][0]

    def test_degrades_when_pptx_missing(self):
        """Returns empty list without raising if python-pptx is not installed."""
        from backend.services.ingest import get_pptx_pages

        with patch.dict(sys.modules, {"pptx": None}):
            result = get_pptx_pages("deck.pptx")

        assert result == []


# ─────────────────────────────────────────────────────────────────────────────
# XLSX
# ─────────────────────────────────────────────────────────────────────────────

class TestGetXlsxChunks:
    def _write_xlsx(self, sheet_data: dict) -> str:
        """Write a minimal XLSX file using openpyxl and return its path."""
        import openpyxl
        wb = openpyxl.Workbook()
        first = True
        for sheet_name, rows in sheet_data.items():
            if first:
                ws = wb.active
                ws.title = sheet_name
                first = False
            else:
                ws = wb.create_sheet(sheet_name)
            for row in rows:
                ws.append(row)
        path = tempfile.mktemp(suffix=".xlsx")
        wb.save(path)
        return path

    def test_every_chunk_has_sheet_and_header(self):
        from backend.services.ingest import get_xlsx_chunks, _CSV_ROWS_PER_CHUNK

        rows = [["id", "name"]] + [[str(i), f"item_{i}"] for i in range(_CSV_ROWS_PER_CHUNK + 3)]
        path = self._write_xlsx({"Data": rows})
        try:
            chunks = get_xlsx_chunks(path)
        finally:
            os.unlink(path)

        assert len(chunks) == 2
        for text, _ in chunks:
            first_line = text.splitlines()[0]
            assert "[Sheet: Data]" in first_line
            assert "id" in first_line
            assert "name" in first_line

    def test_multiple_sheets_all_indexed(self):
        from backend.services.ingest import get_xlsx_chunks

        path = self._write_xlsx({
            "Sheet1": [["a", "b"], [1, 2]],
            "Sheet2": [["x", "y"], [3, 4]],
        })
        try:
            chunks = get_xlsx_chunks(path)
        finally:
            os.unlink(path)

        texts = " ".join(c[0] for c in chunks)
        assert "Sheet1" in texts
        assert "Sheet2" in texts

    def test_header_only_sheet_returns_one_chunk(self):
        from backend.services.ingest import get_xlsx_chunks

        path = self._write_xlsx({"Empty": [["col1", "col2"]]})
        try:
            chunks = get_xlsx_chunks(path)
        finally:
            os.unlink(path)

        assert len(chunks) == 1
        assert "col1" in chunks[0][0]

    def test_chunk_numbers_are_sequential(self):
        from backend.services.ingest import get_xlsx_chunks

        path = self._write_xlsx({
            "S1": [["a"], [1], [2]],
            "S2": [["b"], [3], [4]],
        })
        try:
            chunks = get_xlsx_chunks(path)
        finally:
            os.unlink(path)

        nums = [c[1] for c in chunks]
        assert nums == list(range(1, len(chunks) + 1))

    def test_degrades_when_openpyxl_missing(self):
        from backend.services.ingest import get_xlsx_chunks

        with patch.dict(sys.modules, {"openpyxl": None}):
            result = get_xlsx_chunks("data.xlsx")

        assert result == []


# ─────────────────────────────────────────────────────────────────────────────
# HTML
# ─────────────────────────────────────────────────────────────────────────────

class TestGetHtmlPages:
    def _write_html(self, content: str) -> str:
        f = tempfile.NamedTemporaryFile(
            mode="w", suffix=".html", delete=False, encoding="utf-8"
        )
        f.write(content)
        f.flush()
        return f.name

    def test_trafilatura_path_used_when_available(self):
        from backend.services.ingest import get_html_pages

        html = "<html><body><p>Hello world</p></body></html>"
        path = self._write_html(html)
        try:
            fake_trafilatura = MagicMock()
            fake_trafilatura.extract.return_value = "Hello world"
            with patch.dict(sys.modules, {"trafilatura": fake_trafilatura}):
                pages = get_html_pages(path)
        finally:
            os.unlink(path)

        assert len(pages) == 1
        assert pages[0][0] == "Hello world"
        assert pages[0][1] == 1

    def test_fallback_to_stdlib_when_trafilatura_fails(self):
        from backend.services.ingest import get_html_pages

        html = "<html><body><p>Fallback content</p></body></html>"
        path = self._write_html(html)
        try:
            fake_trafilatura = MagicMock()
            fake_trafilatura.extract.side_effect = Exception("boom")
            with patch.dict(sys.modules, {"trafilatura": fake_trafilatura}):
                pages = get_html_pages(path)
        finally:
            os.unlink(path)

        assert len(pages) == 1
        assert "Fallback content" in pages[0][0]

    def test_stdlib_strips_script_and_style(self):
        from backend.services.ingest import get_html_pages

        html = """<html>
        <head><style>body { color: red; }</style></head>
        <body>
          <script>alert('hi')</script>
          <p>Visible text</p>
        </body></html>"""
        path = self._write_html(html)
        try:
            fake_trafilatura = MagicMock()
            fake_trafilatura.extract.return_value = None  # force stdlib fallback
            with patch.dict(sys.modules, {"trafilatura": fake_trafilatura}):
                pages = get_html_pages(path)
        finally:
            os.unlink(path)

        assert len(pages) == 1
        text = pages[0][0]
        assert "Visible text" in text
        assert "alert" not in text
        assert "color: red" not in text

    def test_empty_html_returns_empty_list(self):
        from backend.services.ingest import get_html_pages

        path = self._write_html("")
        try:
            result = get_html_pages(path)
        finally:
            os.unlink(path)

        assert result == []


# ─────────────────────────────────────────────────────────────────────────────
# process_single_document doc_type routing
# ─────────────────────────────────────────────────────────────────────────────

class TestNewFormatDocTypes:
    def _write(self, content: str, ext: str) -> tuple[str, str]:
        tmp = tempfile.mkdtemp()
        filename = f"testfile{ext}"
        with open(os.path.join(tmp, filename), "w", encoding="utf-8") as f:
            f.write(content)
        return tmp, filename

    def test_html_gets_html_doc_type(self, monkeypatch):
        from backend.services import ingest

        docs_dir, filename = self._write(
            "<html><body><p>Hello</p></body></html>", ".html"
        )
        monkeypatch.setattr(ingest.settings, "docs_dir", docs_dir)
        try:
            docs = ingest.process_single_document(filename)
        finally:
            import shutil; shutil.rmtree(docs_dir)

        assert docs, "Expected at least one chunk"
        assert docs[0]["type"] == "html"

    def test_htm_extension_also_accepted(self, monkeypatch):
        from backend.services import ingest

        docs_dir, filename = self._write(
            "<html><body><p>Hi</p></body></html>", ".htm"
        )
        monkeypatch.setattr(ingest.settings, "docs_dir", docs_dir)
        try:
            docs = ingest.process_single_document(filename)
        finally:
            import shutil; shutil.rmtree(docs_dir)

        assert docs[0]["type"] == "html"


# ─────────────────────────────────────────────────────────────────────────────
# _ALLOWED_EXTENSIONS coverage
# ─────────────────────────────────────────────────────────────────────────────

class TestAllowedExtensions:
    def test_ingest_allowed_extensions_includes_new_formats(self):
        from backend.services.ingest import _ALLOWED_EXTENSIONS

        for ext in (".pptx", ".xlsx", ".html", ".htm"):
            assert ext in _ALLOWED_EXTENSIONS, f"{ext} missing from ingest._ALLOWED_EXTENSIONS"

    def test_router_allowed_extensions_includes_new_formats(self):
        from backend.routers.knowledge import _ALLOWED_EXTENSIONS as router_exts

        for ext in (".pptx", ".xlsx", ".html", ".htm"):
            assert ext in router_exts, f"{ext} missing from router._ALLOWED_EXTENSIONS"

    def test_router_allowed_mime_types_includes_pptx(self):
        from backend.routers.knowledge import _ALLOWED_MIME_TYPES

        assert any("presentationml" in m for m in _ALLOWED_MIME_TYPES)

    def test_router_allowed_mime_types_includes_xlsx(self):
        from backend.routers.knowledge import _ALLOWED_MIME_TYPES

        assert any("spreadsheetml" in m for m in _ALLOWED_MIME_TYPES)
